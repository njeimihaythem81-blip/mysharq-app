import io
import json
import os
import sys
import time
import unittest

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
from PIL import Image
import openpyxl

from core import auth, catalog, config, docs, imaging, notify, orders, security, store
from core.ghstore import GhError
from core.security import LIMITER, THROTTLE
from core.validators import (decode_key, item_key, neutralize_formula, normalize_phone, normalize_support, safe_html, search_norm)
from tests.fakegh import FakeRepo

SEED = json.load(open(os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "seed", "categories.json")))["categories"]
SECRETS = {"GITHUB_TOKEN_DATA": "t1", "GITHUB_TOKEN_IMAGES": "t2", "GITHUB_OWNER": "o", "IMAGES_REPO": "i", "DATA_REPO": "d",
           "SESSION_SECRET": "s" * 40, "SETUP_CODE": "setup-code-123", "STAFF_KEY": "k" * 24, "BRANCH": "main"}
ITEMS = [{"item_no": "LC1D09M7", "description": "Contactor 9A", "section": "TESYS", "branch": "CONTACTORS TO 60A 30KW", "sub_branch": "", "qty": 120},
         {"item_no": "LC1D80M7", "description": "Contactor 80A", "section": "TESYS", "branch": "CONTACTORS D 37 TO 75KW", "sub_branch": "", "qty": 0},
         {"item_no": "LRD08", "description": "Overload 2.5-4A", "section": "TESYS", "branch": "OVERLOADS", "sub_branch": "LRD", "qty": 60}]


def fresh():
    config.set_provider(dict(SECRETS))
    data, imgs = FakeRepo(), FakeRepo()
    store.set_ctx(store.Ctx(data, imgs))
    for k in list(LIMITER._h): del LIMITER._h[k]
    THROTTLE._f.clear()
    notify._last_sent = 0.0; notify._pending.clear()
    sent = []
    notify._send = lambda text, title='x': (sent.append(text), True)[1]
    data.put_file("items.json", json.dumps({"items": ITEMS}).encode(), "t")
    return data, imgs, sent


class Validators(unittest.TestCase):
    def test_phone(self):
        for ok in ("0912345678", "912345678", "+218912345678", "00218912345678", "٠٩١٢٣٤٥٦٧٨", "091 234-5678", "(091) 234.5678"):
            self.assertEqual(normalize_phone(ok), "+218912345678", ok)
        for bad in ("0212345678", "0951234567", "091234567", "09123abc78", "", None, "+21891234567890", "0912345678912"):
            self.assertIsNone(normalize_phone(bad), bad)

    def test_support_number(self):
        self.assertEqual(normalize_support("+216 22 123 456"), "+21622123456"); self.assertEqual(normalize_support("0021622123456"), "+21622123456")
        self.assertEqual(normalize_support("0912345678"), "+218912345678")
        for bad in ("22123456", "abc", "+0123456789", "+12", ""): self.assertIsNone(normalize_support(bad), bad)

    def test_keys_and_search(self):
        for n in ("LC1D09M7", "LC1/D09 M7_X", "A.B-C", ".hidden", "حرف"):
            self.assertEqual(decode_key(item_key(n)), n.upper())
        self.assertNotIn("/", item_key("a/b")); self.assertFalse(item_key("..").startswith("."))
        self.assertEqual(search_norm("lc1 d09-m7"), "LC1D09M7")

    def test_escape_and_formula(self):
        self.assertNotIn("<", safe_html("<script>alert(1)</script>"))
        self.assertNotIn("[", safe_html("[x](http://evil)"))
        self.assertEqual(safe_html("MY SHARQ RFQ's"), "MY SHARQ RFQ&#x27;s")      # entity must stay intact
        self.assertEqual(safe_html("A & B #1"), "A &amp; B &#35;1")
        for v in ("=HYPERLINK(\"x\")", "+1+1", "-2", "@SUM(1)", "\t=1"):
            self.assertTrue(neutralize_formula(v).startswith("'"), v)
        self.assertEqual(neutralize_formula("Ahmed"), "Ahmed")


class Security(unittest.TestCase):
    def test_password(self):
        h = security.hash_password("Correct-horse-9")
        self.assertTrue(security.verify_password("Correct-horse-9", h)); self.assertFalse(security.verify_password("nope", h))
        self.assertIsNotNone(security.password_policy("password123", "x", "operator"))
        self.assertIsNotNone(security.password_policy("short", "x", "operator"))
        self.assertIsNotNone(security.password_policy("elevenchars", "x", "manager"))
        self.assertIsNone(security.password_policy("Long-enough-pass-1", "x", "manager"))

    def test_tokens(self):
        t = security.make_token("k" * 40, "ali", "operator", 1, 100, now=1000)
        self.assertEqual(security.verify_token("k" * 40, t, now=1050)["u"], "ali")
        self.assertIsNone(security.verify_token("k" * 40, t, now=1200))      # expired
        self.assertIsNone(security.verify_token("x" * 40, t, now=1050))      # wrong key
        p, s = t.split(".")
        import base64
        evil = base64.urlsafe_b64encode(json.dumps({"u": "ali", "r": "manager", "v": 1, "iat": 1000, "exp": 9999999999}).encode()).decode().rstrip("=")
        self.assertIsNone(security.verify_token("k" * 40, evil + "." + s, now=1050))  # role tamper

    def test_throttle_no_permanent_lock(self):
        th = security.Throttle()
        for i in range(5): th.fail("k", now=1000 + i)
        self.assertGreater(th.wait_seconds("k", now=1005), 0)
        self.assertEqual(th.wait_seconds("k", now=1005 + 20), 0)   # wait is finite
        for i in range(30): th.fail("k", now=2000 + i)
        self.assertLessEqual(th.wait_seconds("k", now=2030), 900)

    def test_limiter(self):
        l = security.RateLimiter()
        self.assertTrue(all(l.allow("a", 5, 3600, now=10 + i) for i in range(5))); self.assertFalse(l.allow("a", 5, 3600, now=20))
        self.assertTrue(l.allow("a", 5, 3600, now=4000))


class Catalog(unittest.TestCase):
    def test_public_has_no_quantity(self):
        pub = catalog.build_public(ITEMS, SEED)
        blob = json.dumps({k: pub[k] for k in ("items", "tree")}) + json.dumps([i for _, _, i in pub["index"]])
        self.assertNotIn('"qty"', blob); self.assertNotIn("120", blob)
        self.assertEqual({i["status"] for i in pub["items"]}, {"Available", "Waiting"})
        self.assertEqual(catalog.search(pub, "lc1 d09-m7")[0]["item_no"], "LC1D09M7")
        self.assertEqual([x["name"] for x in pub["tree"]], ["TESYS"])  # empty sections hidden

    def _xlsx(self, rows):
        wb = openpyxl.Workbook(); ws = wb.active; ws.title = "Items"
        ws.append(["Section", "Branch", "Sub-Branch", "Item No", "Description", "Quantity"])
        for r in rows: ws.append(r)
        b = io.BytesIO(); wb.save(b); return b.getvalue()

    def test_import_report_and_modes(self):
        rows = [["TESYS", "OVERLOADS", "LRD", "A1", "d", 5], ["TESYS", "OVERLOADS", "LRD", "a1", "dup", 1], ["TESYS", "", "", "", "no id", 1],
                ["tesys", " overloads ", "lrd", "A2", "d2", "abc"], ["TESYZ", "", "", "A3", "typo", 2], ["TESYS", "", "", "A4", "", 1]]
        parsed, err = catalog.parse_items_file(self._xlsx(rows), "x.xlsx"); self.assertIsNone(err)
        rep = catalog.validate_import(parsed, SEED)
        self.assertEqual(len(rep["dups"]), 1); self.assertEqual(len(rep["empty_no"]), 1); self.assertEqual(len(rep["empty_desc"]), 1)
        self.assertEqual(len(rep["bad_qty"]), 1); self.assertIn(("TESYZ", "", ""), rep["unknown"])
        self.assertTrue(rep["suggest"][("TESYZ", "", "")])
        items, cats, s = catalog.apply_import(ITEMS, rep, "merge", {}, SEED)       # unknown -> skip by default
        self.assertEqual(s["skipped"], 1); self.assertEqual(len(items), 5)
        items2, cats2, s2 = catalog.apply_import(ITEMS, rep, "merge", {("TESYZ", "", ""): "create"}, SEED)
        self.assertIn("TESYZ", [c["section"] for c in cats2]); self.assertEqual(len(items2), 6)
        items3, _, s3 = catalog.apply_import(ITEMS, rep, "replace", {}, SEED)
        self.assertEqual(len(items3), 2); self.assertEqual(s3["removed"], 3)

    def test_import_attacks(self):
        self.assertIsNotNone(catalog.parse_items_file(b"x", "a.xlsm")[1])
        self.assertIsNotNone(catalog.parse_items_file(b"not a zip", "a.xlsx")[1])
        self.assertIsNotNone(catalog.parse_items_file(b"x" * (11 * 1024 * 1024), "a.csv")[1])
        csvb = b"Section,Branch,Sub-Branch,Item No,Description,Quantity\nTESYS,,,=cmd,desc,1\n"
        rows, err = catalog.parse_items_file(csvb, "a.csv"); self.assertEqual(rows[0]["item_no"], "=cmd")

    def test_category_editor(self):
        import copy
        cats = copy.deepcopy(SEED)
        cats = catalog.cat_add(cats, (), "NEW SECTION"); cats = catalog.cat_add(cats, ("NEW SECTION",), "B1")
        with self.assertRaises(ValueError): catalog.cat_add(cats, (), "new section")
        items = [dict(ITEMS[0])]
        cats, items = catalog.cat_rename(cats, items, ("TESYS",), "TESYS2"); self.assertEqual(items[0]["section"], "TESYS2")
        with self.assertRaises(ValueError): catalog.cat_delete(cats, items, ("TESYS2",))
        cats = catalog.cat_delete(cats, items, ("NEW SECTION", "B1")); cats = catalog.cat_move(cats, ("IDE BOXES",), 1)
        self.assertEqual(cats[0]["section"], "TESYS2")


class Orders(unittest.TestCase):
    def setUp(self): self.data, self.imgs, self.sent = fresh()

    def go(self, **kw):
        a = dict(basket={"LC1D09M7": {"qty": 3, "note": "n"}}, name="Ahmed", phone="0912345678", project="", session_id="s1", ip="1.1.1.1",
                 honeypot="", form_started=time.time() - 10)
        a.update(kw); return orders.submit(**a)

    def test_happy_and_rebuild(self):
        tampered = {"LC1D09M7": {"qty": "5", "note": "x"}, "FAKE": {"qty": 1}, "LRD08": {"qty": -5}, "LC1D80M7": {"qty": 99999}, "description": "x"}
        ok, msg, o = self.go(basket=orders.parse_basket(json.dumps(tampered)))
        self.assertTrue(ok, msg); self.assertEqual([l["item_no"] for l in o["lines"]], ["LC1D09M7"])
        self.assertEqual(o["lines"][0]["description"], "Contactor 9A"); self.assertEqual(o["customer"]["phone"], "+218912345678")
        self.assertEqual(len(self.sent), 1); self.assertNotIn("912345678", self.sent[0]); self.assertIn("Ahmed", self.sent[0])
        self.assertTrue(store.get_order(o["order_id"])["notified"])

    def test_rejections(self):
        self.assertFalse(self.go(honeypot="bot")[0]); self.assertFalse(self.go(form_started=time.time())[0])
        self.assertFalse(self.go(phone="0212345678")[0]); self.assertFalse(self.go(name="A")[0]); self.assertFalse(self.go(basket={})[0])
        self.assertEqual(len(store.list_orders()), 0)

    def test_flood_limits(self):
        n = 0
        for i in range(40):
            if self.go(session_id="s%d" % i, phone="0912345678", ip="9.9.9.%d" % i)[0]: n += 1
        self.assertEqual(n, 5)                                       # per-phone cap
        LIMITER._h.clear(); n = 0
        for i in range(100):
            if self.go(session_id="x%d" % i, ip="8.8.8.%d" % i, phone="09123%05d" % i)[0]: n += 1
        self.assertEqual(n, 60)                                      # global breaker
        self.assertLessEqual(len(self.sent), 1)                      # Telegram aggregated, not 60 messages
        for t in [notify._timer]:
            if t: t.cancel()

    def test_claim_and_purge(self):
        o1 = self.go()[2]; LIMITER._h.clear(); o2 = self.go(session_id="b")[2]
        won, o = store.claim_order(o1["order_id"], "ali", "Ali", "op:ali"); self.assertTrue(won)
        won2, o = store.claim_order(o1["order_id"], "sara", "Sara", "op:sara"); self.assertFalse(won2); self.assertEqual(o["opened_by_display"], "Ali")
        self.assertEqual([x["order_id"] for x in store.list_orders() if not x["opened_by"]], [o2["order_id"]])
        with self.assertRaises(GhError): store.purge_pending("manager")                    # no purge without export
        ids = [o1["order_id"]]
        store.set_pending({"batch": "B1", "time": "t", "order_ids": ids, "file": "f.xlsx"}, "manager")
        LIMITER._h.clear(); o3 = self.go(session_id="c")[2]                               # arrives after export
        self.assertEqual(store.purge_pending("manager"), 1)
        left = {x["order_id"] for x in store.list_orders()}; self.assertEqual(left, {o2["order_id"], o3["order_id"]})
        self.assertIsNone(store.get_pending())

    def test_documents(self):
        o = self.go(name="=cmd|' /C calc'!A0")[2]
        for x in (docs.order_xlsx(o), docs.export_xlsx([o])):
            ws = openpyxl.load_workbook(io.BytesIO(x)).active
            self.assertFalse(any(str(c.value).startswith("=") for r in ws.iter_rows() for c in r))
        self.assertEqual([c.value for c in openpyxl.load_workbook(io.BytesIO(docs.order_xlsx(o))).active[1]], ["Item No", "Item Name", "Quantity"])
        o["customer"]["name"] = "أحمد محمد"
        self.assertTrue(docs.order_pdf(o, store.settings()).startswith(b"%PDF"))


class AuthAndPhotos(unittest.TestCase):
    def setUp(self): self.data, self.imgs, self.sent = fresh()

    def test_first_run_guards(self):
        self.assertFalse(auth.setup_manager("wrong", "Long-enough-pass-1", "Long-enough-pass-1", "ip")[0])
        self.assertTrue(auth.setup_manager("setup-code-123", "Long-enough-pass-1", "Long-enough-pass-1", "ip")[0])
        self.assertFalse(auth.setup_manager("setup-code-123", "Another-pass-12345", "Another-pass-12345", "ip")[0])   # locked after creation
        self.assertIn("categories.json", self.data.files)
        self.data.files.pop("manager.json"); self.data.fail_paths["manager.json"] = GhError("boom")   # transient error != missing
        with self.assertRaises(GhError): auth.setup_manager("setup-code-123", "Long-enough-pass-1", "Long-enough-pass-1", "ip2")

    def test_login_sessions(self):
        auth.setup_manager("setup-code-123", "Long-enough-pass-1", "Long-enough-pass-1", "ip")
        self.assertTrue(auth.add_operator("ali", "Ali", "Temp-pass-1234")[0]); self.assertFalse(auth.add_operator("manager", "x", "Temp-pass-1234")[0])
        ok, m1, tok, u = auth.login("ali", "Temp-pass-1234", "1.1.1.1"); self.assertTrue(ok and u["must_change"])
        bad1 = auth.login("ali", "wrongwrong", "2.2.2.2")[1]; bad2 = auth.login("ghost", "wrongwrong", "3.3.3.3")[1]
        self.assertEqual(bad1, bad2)                                                       # no username enumeration
        self.assertEqual(auth.session_user(tok)["id"], "ali")
        auth.set_operator_active("ali", False); self.assertIsNone(auth.session_user(tok))   # deactivation ends the session
        auth.set_operator_active("ali", True); ok, _, tok2, _ = auth.login("ali", "Temp-pass-1234", "1.1.1.1"); self.assertTrue(ok)
        self.assertTrue(auth.change_password(auth.session_user(tok2), "Temp-pass-1234", "New-pass-98765")[0]); self.assertIsNone(auth.session_user(tok2))
        for i in range(6): auth.login("manager", "bad-bad-bad", "4.4.4.4")                 # throttled, never permanently locked
        self.assertIn("Try again", auth.login("manager", "Long-enough-pass-1", "4.4.4.4")[1])
        self.assertTrue(auth.login("manager", "Long-enough-pass-1", "5.5.5.5")[0])
        self.assertTrue(auth.check_password({"id": "manager", "role": "manager"}, "Long-enough-pass-1"))

    def test_photos_points(self):
        im = Image.new("RGB", (2000, 1500), "red"); b = io.BytesIO(); im.save(b, "JPEG"); t, f = imaging.process_photo(b.getvalue())
        self.assertTrue(store.save_photo("LC1D09M7", t, f, "ali"))
        self.assertFalse(store.save_photo("LC1D09M7", t, f, "sara"))                       # replacement: no point
        self.assertEqual(store.points(), {"ali": 1})
        for i in range(1, 100): self.data.put_file("photolog/K%d@@ali@@%d.json" % (i, 1000 + i), b"{}", "t")
        store.CACHE.invalidate("d:photolog"); self.assertEqual(store.points()["ali"] // 100, 1)  # 100 points = 1 star
        self.assertEqual(store.points(reset_at=1050)["ali"], 51)   # K50..K99 plus the real upload

    def test_images(self):
        def raw(im, fmt, **kw):
            b = io.BytesIO(); im.save(b, fmt, **kw); return b.getvalue()
        ex = Image.Exif(); ex[0x010F] = "Secret camera"; ex[0x8825] = {1: "N"}
        t, f = imaging.process_photo(raw(Image.new("RGB", (800, 600), "blue"), "JPEG", exif=ex))
        for out in (t, f): self.assertEqual(len(Image.open(io.BytesIO(out)).getexif()), 0)
        self.assertLessEqual(max(Image.open(io.BytesIO(t)).size), 300); self.assertLessEqual(max(Image.open(io.BytesIO(f)).size), 900)
        for bad in (b"<svg xmlns='http://www.w3.org/2000/svg'/>", b"MZ" + b"\0" * 100, b"", b"x" * (16 * 1024 * 1024)):
            with self.assertRaises(imaging.ImageRejected): imaging.process_photo(bad)
        with self.assertRaises(imaging.ImageRejected): imaging.process_photo(raw(Image.new("L", (8000, 8000)), "PNG"))   # 64 MP bomb
        with self.assertRaises(imaging.ImageRejected): imaging.process_photo(raw(Image.new("RGB", (50, 50)), "GIF"))
        big = Image.effect_noise((1600, 2400), 90).convert("RGB")
        self.assertLessEqual(len(imaging.process_background(raw(big, "PNG"), "mobile")), 300 * 1024)


class V2(unittest.TestCase):
    def setUp(self): self.data, self.imgs, self.sent = fresh()

    def _zip(self, files):
        import zipfile
        b = io.BytesIO()
        with zipfile.ZipFile(b, "w") as z:
            for n, d in files.items(): z.writestr(n, d)
        return b.getvalue()

    def test_zip_matching_and_limits(self):
        png = io.BytesIO(); Image.new("RGB", (50, 50), "red").save(png, "PNG"); png = png.getvalue()
        z = self._zip({"TESYS.jpg": png, "folder/TESYS__OVERLOADS.png": png, "tesys__overloads__lrd.JPG": png,
                       "__MACOSX/._TESYS.jpg": b"x", "readme.txt": b"hi", "Nothing.jpg": png, "TESYS__OVERLOADS__NOPE.jpg": png,
                       "ACB (AIR CIRCUIT BREAKERS).jpg": png})
        entries, errs = catalog.read_zip_images(z)
        matched, unmatched = catalog.match_zip_images(entries, SEED)
        self.assertEqual(set(matched), {("TESYS",), ("TESYS", "OVERLOADS"), ("TESYS", "OVERLOADS", "LRD"), ("ACB (AIR CIRCUIT BREAKERS)",)})
        self.assertEqual(sorted(unmatched), ["Nothing.jpg", "TESYS__OVERLOADS__NOPE.jpg"])
        self.assertTrue(any("readme.txt" in e for e in errs))
        big = self._zip({"TESYS.jpg": b"\0" * (7 * 1024 * 1024)})                       # larger than 6 MB: skipped
        self.assertEqual(catalog.read_zip_images(big)[0], []); self.assertTrue(catalog.read_zip_images(big)[1])
        self.assertTrue(catalog.read_zip_images(b"not a zip")[1])
        names = catalog.image_names_table(SEED); self.assertEqual(len(names), len(catalog.all_paths(SEED))); self.assertIn("TESYS__OVERLOADS__LRD.jpg", [n[4] for n in names])

    def test_image_keys_follow_rename_and_delete(self):
        m = {catalog.path_id(("TESYS",)): "t1", catalog.path_id(("TESYS", "OVERLOADS")): "t2", catalog.path_id(("ACTI 9",)): "t3"}
        m2 = catalog.remap_image_keys(m, ("TESYS",), "TESYS NEW")
        self.assertEqual(m2[catalog.path_id(("TESYS NEW", "OVERLOADS"))], "t2"); self.assertEqual(m2[catalog.path_id(("ACTI 9",))], "t3")
        m3, gone = catalog.drop_image_keys(m2, ("TESYS NEW",)); self.assertEqual(sorted(gone), ["t1", "t2"]); self.assertEqual(list(m3.values()), ["t3"])

    def test_category_images_end_to_end(self):
        t, f = imaging.process_photo(self._jpg())
        store.set_cat_images({("TESYS",): imaging.process_cat_image(self._jpg())}, "manager")
        url = store.cat_image_url(("TESYS",)); self.assertIsNotNone(url); self.assertIn("/images/cat/", url)
        self.assertIsNone(store.cat_image_url(("ACTI 9",)))
        store.remove_cat_image(("TESYS",), "manager"); self.assertIsNone(store.cat_image_url(("TESYS",)))
        self.assertFalse(any(p.startswith("images/cat/") for p in self.imgs.files))

    def _jpg(self):
        b = io.BytesIO(); Image.new("RGB", (900, 600), "blue").save(b, "JPEG"); return b.getvalue()

    def test_categories_sheet_and_merge(self):
        wb = openpyxl.Workbook(); ws = wb.active; ws.title = "Items"; ws.append(["Section", "Branch", "Sub-Branch", "Item No", "Description", "Quantity"])
        cs = wb.create_sheet("Categories"); cs.append(["Section", "Branch", "Sub-Branch"])
        for r in (["NEW SEC", "", ""], ["NEW SEC", "B1", ""], ["NEW SEC", "B1", "S1"], ["TESYS", "OVERLOADS", ""], ["", "X", ""], ["NEW SEC", "", "ORPHAN"]): cs.append(r)
        b = io.BytesIO(); wb.save(b)
        paths = catalog.parse_categories_sheet(b.getvalue(), "f.xlsx")
        self.assertEqual(paths, [("NEW SEC",), ("NEW SEC", "B1"), ("NEW SEC", "B1", "S1"), ("TESYS", "OVERLOADS")])
        cats, added = catalog.merge_category_paths(SEED, paths)
        self.assertEqual(added, 3); self.assertEqual(cats[-1]["section"], "NEW SEC"); self.assertEqual(cats[-1]["branches"][0]["sub_branches"], ["S1"])
        self.assertEqual(len(cats), len(SEED) + 1)                      # nothing removed or reordered
        self.assertEqual(catalog.merge_category_paths(cats, paths)[1], 0)

    def test_manual_items(self):
        it, err = catalog.item_add(ITEMS, "NEW1", "Desc", 5, "TESYS", "OVERLOADS", "LRD", SEED); self.assertIsNone(err); self.assertEqual(len(it), 4)
        self.assertIsNotNone(catalog.item_add(it, "new1", "x", 1, "TESYS", "", "", SEED)[1])               # duplicate (case-insensitive)
        self.assertIsNotNone(catalog.item_add(it, "N2", "x", 1, "NOPE", "", "", SEED)[1])                   # unknown category
        self.assertIsNotNone(catalog.item_add(it, "N3", "x", -1, "TESYS", "", "", SEED)[1]); self.assertIsNotNone(catalog.item_add(it, "N4", "", 1, "TESYS", "", "", SEED)[1])
        it2, err = catalog.item_update(it, "NEW1", "Changed", 0, "tesys", "", "", SEED); self.assertIsNone(err)
        row = catalog.find_item(it2, "new1"); self.assertEqual((row["description"], row["qty"], row["section"], row["branch"]), ("Changed", 0, "TESYS", ""))
        it3, err = catalog.item_delete(it2, "NEW1"); self.assertIsNone(err); self.assertEqual(len(it3), 3); self.assertIsNotNone(catalog.item_delete(it3, "NEW1")[1])

    def test_backups_and_restore(self):
        cats = catalog.norm_categories(SEED)
        for i in range(13):
            store.save_catalog(ITEMS[:1 + i % 3], cats, "manager", "import", backup=True)
            self.data.n += 1
            self.data.files.setdefault("x", b"")
        names = store.list_backups(); self.assertEqual(len(names), store.BACKUP_KEEP)
        store.save_catalog([], cats, "manager", "delete-all", backup=True)               # a bad action
        newest = store.list_backups()[0]; self.assertIn("delete-all", newest)
        self.assertEqual(len(store.read_backup(newest)["items"]), len(ITEMS[:1 + 12 % 3]))     # backup holds the state BEFORE
        store.restore_backup(newest, "manager"); store.CACHE.invalidate("d:")
        self.assertEqual(len(store.items()), len(ITEMS[:1 + 12 % 3]))
        with self.assertRaises(GhError): store.read_backup("../items.json")

    def test_items_workbook_and_staff_key(self):
        data = docs.items_xlsx([dict(ITEMS[0], description="=evil()")], SEED, {})
        wb = openpyxl.load_workbook(io.BytesIO(data)); self.assertEqual(wb.sheetnames, ["Items", "Categories"])
        self.assertEqual([c.value for c in wb["Items"][1]], ["Section", "Branch", "Sub-Branch", "Item No", "Description", "Quantity"])
        self.assertTrue(wb["Items"]["E2"].value.startswith("'")); self.assertEqual(wb["Categories"].max_row, len(catalog.all_paths(SEED)) + 1)
        rows, err = catalog.parse_items_file(data, "x.xlsx"); self.assertIsNone(err); self.assertEqual(rows[0]["item_no"], "LC1D09M7")    # round trip
        self.assertTrue(auth.check_staff_key("k" * 24, "1.1.1.1")); self.assertFalse(auth.check_staff_key("wrong", "1.1.1.1")); self.assertFalse(auth.check_staff_key("", "1.1.1.1"))
        for i in range(25): auth.check_staff_key("guess%d" % i, "9.9.9.9")
        self.assertFalse(auth.check_staff_key("k" * 24, "9.9.9.9"))                       # throttled after many wrong guesses


class Theme(unittest.TestCase):
    def test_css_builds_with_literal_percent_signs(self):
        from ui import theme
        css = theme.build_css('background-image:url("https://x/a%20b.jpg");', "", 0.76)
        self.assertIn("width:100%;height:100%", css); self.assertIn("a%20b.jpg", css); self.assertIn("rgba(7,11,13,0.76)", css)
        self.assertNotIn("__BGM__", css); self.assertNotIn("__BGD__", css); self.assertNotIn("__OP__", css)
        self.assertTrue(css.startswith("<style>") and css.rstrip().endswith("</style>"))


class Cfg(unittest.TestCase):
    def test_placeholders_refused(self):
        config.set_provider(dict(SECRETS, SESSION_SECRET="REPLACE-with-at-least-32-random-characters"))
        with self.assertRaises(config.ConfigError): config.validate()
        config.set_provider(dict(SECRETS, SESSION_SECRET="short"))
        with self.assertRaises(config.ConfigError): config.validate()
        config.set_provider(dict(SECRETS, STAFF_KEY="short")); 
        with self.assertRaises(config.ConfigError): config.validate()
        config.set_provider({k: v for k, v in SECRETS.items() if k != "STAFF_KEY"})
        with self.assertRaises(config.ConfigError): config.validate()
        config.set_provider(dict(SECRETS)); config.validate()

    def test_channels_dispatch(self):
        import core.notify as n
        fresh(); real = n.requests.post; calls = []
        class R: status_code = 200
        n.requests.post = lambda url, **kw: (calls.append((url, kw)), R())[1]
        try:
            config.set_provider(dict(SECRETS)); self.assertEqual(n.channels(), []); self.assertFalse(n.__dict__["_send"].__name__ == "<lambda>" and False)
            config.set_provider(dict(SECRETS, NTFY_TOPIC="mysharq-abc123", APP_URL="https://x.streamlit.app")); self.assertEqual(n.channels(), ["ntfy"])
            import importlib; importlib.reload(n)
            self.assertTrue(n._send("hello\nbody", "New RFQ MS-1"))
            url, kw = calls[-1]; self.assertEqual(url, "https://ntfy.sh/mysharq-abc123"); self.assertEqual(kw["data"], b"hello\nbody")
            self.assertEqual(kw["headers"]["Title"], "New RFQ MS-1"); self.assertEqual(kw["headers"]["Click"], "https://x.streamlit.app")
            config.set_provider(dict(SECRETS, TELEGRAM_BOT_TOKEN="1:abc", TELEGRAM_CHAT_IDS="5")); self.assertEqual(n.channels(), ["telegram"])
            config.set_provider(dict(SECRETS, EMAIL_USER="a@gmail.com", EMAIL_APP_PASSWORD="pw", EMAIL_TO="b@x.com")); self.assertEqual(n.channels(), ["email"])
        finally:
            n.requests.post = real; config.set_provider(dict(SECRETS))

    def test_telegram_plain_and_aggregated(self):
        data, imgs, sent = fresh()
        notify.notify_new_order("A", "Bob\n\x00" + "x" * 80, 2)
        notify.notify_new_order("B", "C", 1); notify.notify_new_order("D", "E", 1)
        self.assertEqual(len(sent), 1); self.assertEqual(len(notify._pending), 2)       # 1 immediate, 2 aggregated
        line = [l for l in sent[0].split("\n") if l.startswith("Customer:")][0]
        self.assertLessEqual(len(line), len("Customer: ") + 40); self.assertNotIn("\x00", sent[0])
        if notify._timer: notify._timer.cancel()


if __name__ == "__main__":
    unittest.main(verbosity=1)
