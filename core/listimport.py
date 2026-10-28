"""Turn an uploaded list (Excel / CSV / text / text-based PDF) into basket candidates.

Reading happens on the server with small libraries; nothing is stored. Each line is matched with the same smart search
as the search box and given a confidence:  >= AUTO  accepted automatically (customer can still change it),
>= SUGGEST  two suggestions are shown, below that the customer types the item number.
"""
import csv
import io
import re
import zipfile

from . import bundles, catalog
from .validators import clean_text, search_norm, to_western_digits

MAX_BYTES = 5 * 1024 * 1024
MAX_LINES = 150                      # same as the basket limit
MAX_PDF_PAGES = 15
AUTO, SUGGEST = 0.85, 0.45

_X = r"[x\u00d7*]"
_UNITS = r"(?:pcs?|pieces?|units?|ex|qt[e\u00e9]?|\u0642\u0637\u0639\u0647|\u0642\u0637\u0639\u0629|\u062d\u0628\u0647|\u062d\u0628\u0629|\u0639\u062f\u062f)"
_LEAD = re.compile(r"^\s*(\d{1,4})\s*" + _X + r"\s*(?=\S)", re.I)
_TRAIL_X = re.compile(r"(?<=\S)\s*" + _X + r"\s*(\d{1,4})\s*$", re.I)
_TRAIL_SEP = re.compile(r"(?<=\S)\s*[:\-\u2013=]\s*(\d{1,4})\s*" + _UNITS + r"?\s*$", re.I)
_TRAIL_UNIT = re.compile(r"(?<=\S)\s+(\d{1,4})\s*" + _UNITS + r"\s*$", re.I)
_LEAD_UNIT = re.compile(r"^\s*(\d{1,4})\s*" + _UNITS + r"\s+(?=\S)", re.I)
_BULLET = re.compile(r"^\s*(?:[-\u2022*\u2013]+|\(?\d{1,3}[.)\-]\s)\s*")


def parse_line(raw):
    """'2 x contactor 25A' / 'relay 24V : 5' / 'contactor 9A x3' / '5 pcs fuse' -> (text, qty). Quantity is 1..9999, default 1."""
    t = clean_text(to_western_digits(raw), 160)
    t = _BULLET.sub("", t)
    qty = 1
    for rx in (_LEAD, _LEAD_UNIT, _TRAIL_X, _TRAIL_UNIT, _TRAIL_SEP):
        m = rx.search(t)
        if m and 1 <= int(m.group(1)) <= 9999:
            qty = int(m.group(1)); t = (t[:m.start()] + " " + t[m.end():]).strip(); break
    return clean_text(t.strip(" :-\u2013=,;"), 120), qty


# ------------------------------------------------------------------ reading files
def _cell(v):
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return str(v).strip()


def _kind(h):
    h = re.sub(r"[^a-z\u0600-\u06FF]", "", str(h or "").lower())
    if not h:
        return None
    if any(k in h for k in ("qty", "quant", "qte", "\u0643\u0645\u064a", "\u0639\u062f\u062f")):
        return "qty"
    if any(k in h for k in ("itemno", "partno", "code", "ref", "\u0631\u0642\u0645", "\u0643\u0648\u062f", "\u0645\u0631\u062c\u0639", "\u0631\u0645\u0632")):
        return "no"
    if any(k in h for k in ("desc", "design", "name", "item", "\u0648\u0635\u0641", "\u0628\u064a\u0627\u0646", "\u0627\u0633\u0645", "\u0627\u0644\u0635\u0646\u0641", "\u0627\u0644\u0645\u0627\u062f\u0647", "\u0627\u0644\u0645\u0627\u062f\u0629")):
        return "desc"
    return None


def _int_cell(c):
    c = to_western_digits(_cell(c))
    return int(c) if re.fullmatch(r"\d{1,4}", c) and 1 <= int(c) <= 9999 else None


def _rows_to_lines(rows):
    """rows: list of cell lists. Uses a header row when there is one, otherwise guesses (last small whole number = quantity)."""
    hdr, cols = None, {}
    for ri, r in enumerate(rows[:5]):
        k = {}
        for ci, c in enumerate(r):
            kd = _kind(c)
            if kd and kd not in k.values() and not _int_cell(c):
                k[ci] = kd
        if len(k) >= 2 or (len(k) == 1 and ri == 0 and any(v in ("no", "desc") for v in k.values())):
            hdr, cols = ri, {v: ci for ci, v in k.items()}
            break
    out = []
    for r in rows[(hdr + 1) if hdr is not None else 0:]:
        cells = [_cell(c) for c in r]
        if not any(cells):
            continue
        if hdr is not None:
            g = lambda n: cells[cols[n]] if n in cols and cols[n] < len(cells) else ""
            q = _int_cell(g("qty")) or 1
            main, alt = g("no"), g("desc")
            if not main:
                main, alt = alt, ""
            if not main:
                continue
            t, q2 = parse_line(main)
            out.append({"text": t or clean_text(main, 120), "alt": clean_text(alt, 120), "qty": q if "qty" in cols else q2})
        else:
            qty, rest = None, list(cells)
            for i in range(len(rest) - 1, -1, -1):
                v = _int_cell(rest[i])
                if v is not None and any(x for j, x in enumerate(rest) if j != i and x and _int_cell(x) is None):
                    qty = v; rest.pop(i); break
            txt = " ".join(x for x in rest if x)
            t, q2 = parse_line(txt)
            if t:
                out.append({"text": t, "alt": "", "qty": qty or q2})
    return out


def extract_lines(data, filename):
    """-> (lines, notes, error). lines = [{'text', 'alt', 'qty'}] (at most MAX_LINES). notes = informational strings."""
    notes = []
    if len(data) > MAX_BYTES:
        return [], notes, "The file is larger than 5 MB."
    name = (filename or "").lower()
    try:
        if name.endswith(".xlsx"):
            z = zipfile.ZipFile(io.BytesIO(data))
            if any("vbaproject" in n.lower() for n in z.namelist()):
                return [], notes, "Macro-enabled files are not accepted."
            from openpyxl import load_workbook
            wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
            rows = []
            for r in wb.worksheets[0].iter_rows(values_only=True):
                rows.append(list(r))
                if len(rows) > 2000:
                    break
            lines = _rows_to_lines(rows)
        elif name.endswith(".csv"):
            txt = data.decode("utf-8-sig", errors="replace")
            lines = _rows_to_lines([r for _, r in zip(range(2000), csv.reader(io.StringIO(txt)))])
        elif name.endswith(".txt"):
            lines = _text_lines(data.decode("utf-8-sig", errors="replace"))
        elif name.endswith(".pdf"):
            txt, pages = _pdf_text(data)
            if pages > MAX_PDF_PAGES:
                notes.append("Only the first %d pages were read." % MAX_PDF_PAGES)
            if not txt.strip():
                return [], notes, "This PDF has no readable text (it looks scanned). Reading pictures and handwriting is coming in the next version; for now please upload an Excel file or a text-based PDF."
            lines = _text_lines(txt)
        else:
            return [], notes, "Accepted files: .xlsx, .csv, .txt, .pdf"
    except Exception:
        return [], notes, "The file could not be read. Please check it is a valid file."
    if not lines:
        return [], notes, "No lines were found in the file."
    if len(lines) > MAX_LINES:
        notes.append("Only the first %d lines were read (the basket limit)." % MAX_LINES)
        lines = lines[:MAX_LINES]
    return lines, notes, None


def _text_lines(txt):
    out = []
    for raw in txt.replace("\r", "\n").split("\n"):
        t, q = parse_line(raw)
        if len(t) >= 2 and re.search(r"[0-9A-Za-z\u0600-\u06FF]", t):
            out.append({"text": t, "alt": "", "qty": q})
    return out


def _pdf_text(data):
    from pypdf import PdfReader
    rd = PdfReader(io.BytesIO(data))
    if rd.is_encrypted:
        raise ValueError("encrypted")
    parts = []
    for pg in rd.pages[:MAX_PDF_PAGES]:
        parts.append(pg.extract_text() or "")
    return "\n".join(parts), len(rd.pages)


# ------------------------------------------------------------------ matching
def match_text(cat, text):
    """-> (candidates [item_no, ...] best first, confidence 0..1, learned_hit)."""
    res, mode, info = catalog.smart_search_ex(cat, text, 6)
    if not res:
        return [], 0.0, False
    nos = [r["item_no"] for r in res]
    sc, cv = info["scores"], info["cover"]
    s1 = sc.get(nos[0], 0.0)
    s2 = sc.get(nos[1], 0.0) if len(nos) > 1 else 0.0
    if search_norm(text) == search_norm(nos[0]) and search_norm(text):
        return nos, 1.0, False
    if info["learned"] and s1 >= 2000:
        return nos, 0.97, True
    margin = 1.0 if len(nos) == 1 or s2 <= 0 else max(0.0, min(1.0, (s1 - s2) / s1))
    conf = cv.get(nos[0], 0.0) * (0.45 + 0.55 * margin)
    if mode == "closest":
        conf *= 0.7
    return nos, round(min(conf, 0.99), 3), False


def match_line(cat, line):
    """Try the item-number cell first, then the description; keep the better result.
    -> {'text', 'qty', 'cands': [item_no], 'conf', 'state': 'auto'|'suggest'|'manual'|'bundle', 'bundle': index or None}"""
    best = ([], 0.0, False, line["text"])
    for t in (line["text"], line.get("alt", "")):
        if not t:
            continue
        nos, conf, lr = match_text(cat, t)
        if conf > best[1]:
            best = (nos, conf, lr, t)
    nos, conf, lr, used = best
    state = "auto" if conf >= AUTO else ("suggest" if conf >= SUGGEST and nos else "manual")
    out = {"text": line["text"] if line["text"] else used, "wording": used, "qty": line["qty"], "cands": nos[:6], "conf": conf, "state": state, "bundle": None}
    if state != "auto":                                        # a bundle name written in the line beats a doubtful single item
        for t in (line["text"], line.get("alt", "")):
            hit = bundles.find(cat, t, 1) if t else []
            if hit:
                out["state"], out["bundle"] = "bundle", hit[0][0]; break
    return out


def match_all(cat, lines):
    return [dict(match_line(cat, l), id=i) for i, l in enumerate(lines)]
