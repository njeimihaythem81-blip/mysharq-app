"""Catalogue logic: public view (NO quantities), search, Excel import, category editing."""
import csv
import difflib
import io
import re
import zipfile

from .validators import cat_norm, clean_text, item_key, search_norm

MAX_IMPORT_BYTES = 10 * 1024 * 1024
MAX_IMPORT_ROWS = 20000
HEADERS = {"section": "section", "branch": "branch", "subbranch": "sub_branch", "itemno": "item_no",
           "description": "description", "quantity": "qty"}


def availability(qty):
    """Customer-facing status. The number itself must never leave the server."""
    return "Available" if isinstance(qty, int) and qty > 0 else "Waiting"


# ---------------------------------------------------------------- public view
def build_public(items, categories):
    pub, level, counts, by_no, index = [], {}, {}, {}, []
    for it in items:
        d = {"item_no": it["item_no"], "description": it["description"], "section": it["section"],
             "branch": it.get("branch", ""), "sub_branch": it.get("sub_branch", ""),
             "status": availability(it.get("qty", 0)), "key": item_key(it["item_no"])}
        pub.append(d)
        by_no[d["item_no"]] = d
        path = tuple(p for p in (d["section"], d["branch"], d["sub_branch"]) if p)
        level.setdefault(path, []).append(d)
        for i in range(1, len(path) + 1):
            counts[path[:i]] = counts.get(path[:i], 0) + 1
        index.append((search_norm(d["item_no"]), d["description"].lower(), d))
    tree = []
    for c in categories:
        if counts.get((c["section"],), 0) == 0:
            continue
        brs = []
        for b in c.get("branches", []):
            if counts.get((c["section"], b["name"]), 0) == 0:
                continue
            subs = [s for s in b.get("sub_branches", []) if counts.get((c["section"], b["name"], s), 0) > 0]
            brs.append({"name": b["name"], "subs": subs})
        tree.append({"name": c["section"], "branches": brs})
    payload = sorted(([d["item_no"], d["description"][:70], 1 if d["status"] == "Available" else 0] for d in pub),
                     key=lambda r: (search_norm(r[0]), r[0]))
    return {"items": pub, "level": level, "counts": counts, "by_no": by_no, "index": index, "tree": tree, "payload": payload}


def children(cat, path):
    """Visible child category names under path (tuple)."""
    if not path:
        return [s["name"] for s in cat["tree"]]
    for s in cat["tree"]:
        if s["name"] == path[0]:
            if len(path) == 1:
                return [b["name"] for b in s["branches"]]
            for b in s["branches"]:
                if b["name"] == path[1] and len(path) == 2:
                    return list(b["subs"])
    return []


def search(cat, query, limit=500):
    qn = search_norm(query)
    words = [w for w in re.split(r"\s+", query.lower().strip()) if w]
    exact, prefix, contains, desc = [], [], [], []
    for nno, dl, it in cat["index"]:
        if qn and nno == qn:
            exact.append(it)
        elif qn and nno.startswith(qn):
            prefix.append(it)
        elif qn and qn in nno:
            contains.append(it)
        elif words and all(w in dl for w in words):
            desc.append(it)
    return (exact + prefix + contains + desc)[:limit]


# ---------------------------------------------------------------- categories
def norm_categories(cats):
    out = []
    for i, c in enumerate(cats, 1):
        out.append({"order": i, "section": c["section"],
                    "branches": [{"name": b["name"], "sub_branches": list(b.get("sub_branches", []))} for b in c.get("branches", [])]})
    return out


def triples(categories):
    """Canonical (section, branch, sub) strings available in the tree."""
    out = []
    for c in categories:
        out.append((c["section"], "", ""))
        for b in c["branches"]:
            out.append((c["section"], b["name"], ""))
            for s in b["sub_branches"]:
                out.append((c["section"], b["name"], s))
    return out


def _lookup(categories):
    m = {}
    for t in triples(categories):
        m[tuple(cat_norm(x) for x in t)] = t
    return m


def items_using(items, level, name_path):
    """Items referencing a category (name_path = tuple of 1-3 names)."""
    n = len(name_path)
    return [i for i in items if tuple(i.get(k, "") for k in ("section", "branch", "sub_branch"))[:n] == tuple(name_path)]


def cat_add(cats, path, name):
    name = clean_text(name, 80)
    if not name:
        raise ValueError("Name required")
    if len(path) == 0:
        if any(cat_norm(c["section"]) == cat_norm(name) for c in cats):
            raise ValueError("Section already exists")
        cats.append({"order": len(cats) + 1, "section": name, "branches": []})
    else:
        sec = next(c for c in cats if c["section"] == path[0])
        if len(path) == 1:
            if any(cat_norm(b["name"]) == cat_norm(name) for b in sec["branches"]):
                raise ValueError("Branch already exists")
            sec["branches"].append({"name": name, "sub_branches": []})
        else:
            br = next(b for b in sec["branches"] if b["name"] == path[1])
            if any(cat_norm(s) == cat_norm(name) for s in br["sub_branches"]):
                raise ValueError("Sub-branch already exists")
            br["sub_branches"].append(name)
    return norm_categories(cats)


def cat_rename(cats, items, path, new):
    """Rename and update the items that reference it. Returns (cats, items)."""
    new = clean_text(new, 80)
    if not new:
        raise ValueError("Name required")
    key = ("section", "branch", "sub_branch")[len(path) - 1]
    for it in items:
        if tuple(it.get(k, "") for k in ("section", "branch", "sub_branch"))[:len(path)] == tuple(path):
            it[key] = new
    sec = next(c for c in cats if c["section"] == path[0])
    if len(path) == 1:
        sec["section"] = new
    else:
        br = next(b for b in sec["branches"] if b["name"] == path[1])
        if len(path) == 2:
            br["name"] = new
        else:
            br["sub_branches"] = [new if s == path[2] else s for s in br["sub_branches"]]
    return norm_categories(cats), items


def cat_delete(cats, items, path):
    if items_using(items, None, path):
        raise ValueError("This category still has items. Move or delete them first.")
    sec = next(c for c in cats if c["section"] == path[0])
    if len(path) == 1:
        cats = [c for c in cats if c is not sec]
    elif len(path) == 2:
        sec["branches"] = [b for b in sec["branches"] if b["name"] != path[1]]
    else:
        br = next(b for b in sec["branches"] if b["name"] == path[1])
        br["sub_branches"] = [s for s in br["sub_branches"] if s != path[2]]
    return norm_categories(cats)


def cat_move(cats, path, delta):
    def mv(lst, pred):
        i = next(i for i, x in enumerate(lst) if pred(x))
        j = i + delta
        if 0 <= j < len(lst):
            lst[i], lst[j] = lst[j], lst[i]
    if len(path) == 1:
        mv(cats, lambda c: c["section"] == path[0])
    else:
        sec = next(c for c in cats if c["section"] == path[0])
        if len(path) == 2:
            mv(sec["branches"], lambda b: b["name"] == path[1])
        else:
            br = next(b for b in sec["branches"] if b["name"] == path[1])
            mv(br["sub_branches"], lambda s: s == path[2])
    return norm_categories(cats)


# ---------------------------------------------------------------- import
def _hdr(h):
    return HEADERS.get(re.sub(r"[^a-z]", "", str(h or "").lower()))


def _cell(v):
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return str(v).strip()


def parse_items_file(data, filename):
    """-> (rows, fatal_error). Rows keep their spreadsheet row number. Values only, never formulas."""
    if len(data) > MAX_IMPORT_BYTES:
        return [], "File is larger than 10 MB."
    name = (filename or "").lower()
    raw = []
    try:
        if name.endswith(".xlsx"):
            z = zipfile.ZipFile(io.BytesIO(data))
            if any("vbaproject" in n.lower() for n in z.namelist()):
                return [], "Macro-enabled files are not accepted."
            from openpyxl import load_workbook
            wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
            ws = wb["Items"] if "Items" in wb.sheetnames else wb.worksheets[0]
            for i, r in enumerate(ws.iter_rows(values_only=True), 1):
                raw.append((i, list(r)))
                if i > MAX_IMPORT_ROWS + 1:
                    return [], "More than 20,000 rows."
        elif name.endswith(".csv"):
            txt = data.decode("utf-8-sig", errors="replace")
            for i, r in enumerate(csv.reader(io.StringIO(txt)), 1):
                raw.append((i, r))
                if i > MAX_IMPORT_ROWS + 1:
                    return [], "More than 20,000 rows."
        else:
            return [], "Only .xlsx or .csv files are accepted."
    except Exception:
        return [], "The file could not be read. Make sure it is a valid .xlsx or .csv file."
    if not raw:
        return [], "The file is empty."
    cols = {}
    for idx, h in enumerate(raw[0][1]):
        k = _hdr(h)
        if k and k not in cols:
            cols[k] = idx
    missing = [k for k in ("section", "branch", "sub_branch", "item_no", "description", "qty") if k not in cols]
    if missing:
        return [], "Missing columns: " + ", ".join(missing) + ". Expected: Section, Branch, Sub-Branch, Item No, Description, Quantity."
    rows = []
    for i, r in raw[1:]:
        if not any(_cell(c) for c in r):
            continue
        g = lambda k: _cell(r[cols[k]]) if cols[k] < len(r) else ""
        rows.append({"row": i, "section": g("section"), "branch": g("branch"), "sub_branch": g("sub_branch"),
                     "item_no": g("item_no"), "description": g("description"), "qty_raw": g("qty")})
    return rows, None


def validate_import(rows, categories):
    look = _lookup(categories)
    names = [" > ".join(x for x in t if x) for t in triples(categories)]
    rep = {"ok": [], "empty_no": [], "empty_desc": [], "dups": [], "bad_qty": [], "unknown": {}, "suggest": {}}
    seen = set()
    for r in rows:
        no = clean_text(r["item_no"], 80)
        if not no:
            rep["empty_no"].append(r["row"]); continue
        if not clean_text(r["description"], 300):
            rep["empty_desc"].append(r["row"]); continue
        if no.upper() in seen:
            rep["dups"].append((r["row"], no)); continue
        seen.add(no.upper())
        q = r["qty_raw"]
        if q == "":
            qty = 0
        else:
            try:
                f = float(q)
                qty = int(f) if f.is_integer() and f >= 0 else None
            except ValueError:
                qty = None
            if qty is None:
                rep["bad_qty"].append((r["row"], q)); qty = 0
        s, b, sb = (clean_text(r[k], 80) for k in ("section", "branch", "sub_branch"))
        rec = dict(r, item_no=no, description=clean_text(r["description"], 300), qty=qty, section=s, branch=b, sub_branch=sb)
        key = (cat_norm(s), cat_norm(b), cat_norm(sb))
        if (sb and not b) or key not in look:
            raw_t = (s, b, sb)
            rep["unknown"].setdefault(raw_t, []).append(r["row"])
            rec["_unknown"] = raw_t
            if raw_t not in rep["suggest"]:
                label = " > ".join(x for x in raw_t if x)
                close = difflib.get_close_matches(label, names, n=3, cutoff=0.6)
                rep["suggest"][raw_t] = [triples(categories)[names.index(c)] for c in close]
        else:
            rec["section"], rec["branch"], rec["sub_branch"] = look[key]
        rep["ok"].append(rec)
    return rep


def apply_import(existing, report, mode, decisions, categories):
    """decisions: {raw_triple: 'create' | 'skip' | ('map', canonical_triple)}.
    Returns (items, categories, summary). Nothing is lost silently: skipped rows are counted."""
    cats = norm_categories(categories)
    by_no = {i["item_no"].upper(): i for i in existing}
    imported, summary = [], {"added": 0, "updated": 0, "unchanged": 0, "skipped": 0, "removed": 0}
    for rec in report["ok"]:
        t = rec.get("_unknown")
        if t:
            dec = decisions.get(t, "skip")
            if dec == "skip" or t[0] == "" or (t[2] and not t[1]):
                summary["skipped"] += 1; continue
            if isinstance(dec, tuple) and dec[0] == "map":
                rec["section"], rec["branch"], rec["sub_branch"] = dec[1]
            else:
                look = _lookup(cats)
                for lvl, name in enumerate(t):
                    if not name:
                        continue
                    path = tuple(x for x in t[:lvl] if x)
                    if tuple(cat_norm(x) for x in (list(t[:lvl + 1]) + [""] * (2 - lvl))) not in look:
                        cats = cat_add(cats, path, name)
                        look = _lookup(cats)
                canon = look[tuple(cat_norm(x) for x in t)]
                rec["section"], rec["branch"], rec["sub_branch"] = canon
        item = {"item_no": rec["item_no"], "description": rec["description"], "section": rec["section"],
                "branch": rec["branch"], "sub_branch": rec["sub_branch"], "qty": rec["qty"]}
        old = by_no.get(item["item_no"].upper())
        if old is None:
            summary["added"] += 1
        elif all(old.get(k) == item[k] for k in item):
            summary["unchanged"] += 1
        else:
            summary["updated"] += 1
        imported.append(item)
    if mode == "replace":
        new_items = imported
        summary["removed"] = len([k for k in by_no if k not in {i["item_no"].upper() for i in imported}])
    else:
        merged = dict(by_no)
        for it in imported:
            merged[it["item_no"].upper()] = it
        new_items = list(merged.values())
    return new_items, norm_categories(cats), summary


# ======================================================================= v2: images per level, items CRUD, Excel with 2 sheets
import os as _os
import re as _re

SEP = "__"
IMG_EXT = (".jpg", ".jpeg", ".png", ".webp")
ZIP_MAX_FILES, ZIP_MAX_EACH, ZIP_MAX_TOTAL = 400, 6 * 1024 * 1024, 120 * 1024 * 1024


def path_id(path):
    """Stable key of a category path (case/space-insensitive)."""
    return "\x1f".join(cat_norm(p) for p in path)


def all_paths(categories):
    out = []
    for c in categories:
        out.append((c["section"],))
        for b in c.get("branches", []):
            out.append((c["section"], b["name"]))
            for s in b.get("sub_branches", []):
                out.append((c["section"], b["name"], s))
    return out


def _san(name):
    return _re.sub(r'[\\/:*?"<>|]', "-", str(name))


def expected_filename(path):
    return SEP.join(_san(p) for p in path) + ".jpg"


def image_names_table(categories, imgmap=None):
    """Rows for the 'expected image file names' list: Level, Section, Branch, Sub-Branch, File name, Has image."""
    imgmap = imgmap or {}
    rows = []
    for p in all_paths(categories):
        rows.append((("Section", "Branch", "Sub-Branch")[len(p) - 1], p[0], p[1] if len(p) > 1 else "",
                     p[2] if len(p) > 2 else "", expected_filename(p), "yes" if path_id(p) in imgmap else "no"))
    return rows


def read_zip_images(data):
    """-> (entries [(name, bytes)], errors). Hard limits against zip bombs; reads with a size cap."""
    import io as _io
    import zipfile as _zf
    errs, out, total = [], [], 0
    try:
        z = _zf.ZipFile(_io.BytesIO(data))
    except Exception:
        return [], ["The file is not a valid ZIP."]
    infos = z.infolist()
    if len(infos) > ZIP_MAX_FILES * 2:
        return [], ["The ZIP has too many files."]
    for i in infos:
        base = _os.path.basename(i.filename.replace("\\", "/"))
        if i.is_dir() or not base or base.startswith(".") or "__MACOSX" in i.filename:
            continue
        if _os.path.splitext(base)[1].lower() not in IMG_EXT:
            errs.append("Ignored (not an image): " + base); continue
        with z.open(i) as f:
            d = f.read(ZIP_MAX_EACH + 1)
        if len(d) > ZIP_MAX_EACH:
            errs.append("Too large (6 MB max): " + base); continue
        total += len(d)
        if total > ZIP_MAX_TOTAL:
            errs.append("Total size limit reached; the rest was ignored."); break
        out.append((base, d))
        if len(out) >= ZIP_MAX_FILES:
            errs.append("Only the first %d images were read." % ZIP_MAX_FILES); break
    return out, errs


def match_zip_images(entries, categories):
    """-> (matched {path: bytes}, unmatched [names]). Names: SECTION.jpg, SECTION__BRANCH.jpg, SECTION__BRANCH__SUB.jpg."""
    idx = {tuple(cat_norm(_san(p)) for p in path): path for path in all_paths(categories)}
    matched, unmatched = {}, []
    for name, data in entries:
        stem = _os.path.splitext(name)[0]
        key = tuple(cat_norm(x) for x in stem.split(SEP))
        if key in idx:
            matched[idx[key]] = data
        else:
            unmatched.append(name)
    return matched, unmatched


def remap_image_keys(imgmap, old_path, new_name):
    """After renaming the node at old_path, keep its image (and its children's)."""
    pos = len(old_path) - 1
    old_id = path_id(old_path)
    out = {}
    for k, v in imgmap.items():
        if k == old_id or k.startswith(old_id + "\x1f"):
            parts = k.split("\x1f"); parts[pos] = cat_norm(new_name); k = "\x1f".join(parts)
        out[k] = v
    return out


def drop_image_keys(imgmap, path):
    """-> (new_map, removed_tokens) for the node and everything below it."""
    pid = path_id(path)
    keep, gone = {}, []
    for k, v in imgmap.items():
        if k == pid or k.startswith(pid + "\x1f"):
            gone.append(v)
        else:
            keep[k] = v
    return keep, [g for g in gone if g]


# ---- Excel: second sheet "Categories"
def parse_categories_sheet(data, filename):
    """Paths from the optional 'Categories' sheet of an .xlsx (empty list if absent). Values only."""
    import io as _io
    if not (filename or "").lower().endswith(".xlsx"):
        return []
    try:
        from openpyxl import load_workbook
        wb = load_workbook(_io.BytesIO(data), read_only=True, data_only=True)
        if "Categories" not in wb.sheetnames:
            return []
        rows = list(wb["Categories"].iter_rows(values_only=True))
    except Exception:
        return []
    if not rows:
        return []
    cols = {}
    for i, h in enumerate(rows[0]):
        k = _hdr(h)
        if k in ("section", "branch", "sub_branch") and k not in cols:
            cols[k] = i
    if "section" not in cols:
        return []
    out = []
    for r in rows[1:]:
        g = lambda k: clean_text(_cell(r[cols[k]]) if k in cols and cols[k] < len(r) else "", 80)
        s, b, sb = g("section"), g("branch"), g("sub_branch")
        if s and not (sb and not b):
            out.append(tuple(x for x in (s, b, sb) if x))
    return out


def merge_category_paths(categories, paths):
    """Add missing nodes (appended at the end). Never deletes or reorders. -> (categories, added_count)."""
    cats, added = norm_categories(categories), 0
    for path in paths:
        for lvl in range(1, len(path) + 1):
            sub = tuple(path[:lvl])
            existing = {path_id(p) for p in all_paths(cats)}
            if path_id(sub) not in existing:
                cats = cat_add(cats, sub[:-1] if lvl > 1 else (), sub[-1])
                added += 1
    return cats, added


# ---- manual items
def find_item(items, item_no):
    n = clean_text(item_no, 80).upper()
    return next((i for i in items if i["item_no"].upper() == n), None)


def _check_path(categories, section, branch, sub):
    if sub and not branch:
        return "Choose a branch before a sub-branch."
    look = _lookup(categories)
    key = (cat_norm(section), cat_norm(branch), cat_norm(sub))
    return None if key in look else "That Section / Branch / Sub-Branch does not exist."


def item_add(items, item_no, description, qty, section, branch, sub, categories):
    """-> (items, error). Item No is the unique key."""
    no, desc = clean_text(item_no, 80), clean_text(description, 300)
    if not no:
        return items, "Item No is required."
    if not desc:
        return items, "Description is required."
    try:
        q = int(str(qty).strip() or 0)
        assert 0 <= q <= 10_000_000
    except Exception:
        return items, "Quantity must be a whole number (0 or more)."
    if find_item(items, no):
        return items, "This Item No already exists."
    err = _check_path(categories, section, branch, sub)
    if err:
        return items, err
    look = _lookup(categories)
    s, b, sb = look[(cat_norm(section), cat_norm(branch), cat_norm(sub))]
    return items + [{"item_no": no, "description": desc, "section": s, "branch": b, "sub_branch": sb, "qty": q}], None


def item_update(items, item_no, description, qty, section, branch, sub, categories):
    it = find_item(items, item_no)
    if not it:
        return items, "Item not found."
    desc = clean_text(description, 300)
    if not desc:
        return items, "Description is required."
    try:
        q = int(str(qty).strip() or 0)
        assert 0 <= q <= 10_000_000
    except Exception:
        return items, "Quantity must be a whole number (0 or more)."
    err = _check_path(categories, section, branch, sub)
    if err:
        return items, err
    look = _lookup(categories)
    s, b, sb = look[(cat_norm(section), cat_norm(branch), cat_norm(sub))]
    out = [dict(i) for i in items]
    for i in out:
        if i["item_no"].upper() == it["item_no"].upper():
            i.update(description=desc, qty=q, section=s, branch=b, sub_branch=sb)
    return out, None


def item_delete(items, item_no):
    n = clean_text(item_no, 80).upper()
    out = [i for i in items if i["item_no"].upper() != n]
    return out, (None if len(out) < len(items) else "Item not found.")


# ======================================================================= replace the whole tree (from Excel)
def tree_from_paths(paths):
    """Build a category tree from ordered paths (rows of the Categories sheet). Parents are created when missing."""
    cats, _ = merge_category_paths([], paths)
    return cats


def tree_stats(categories):
    secs = len(categories)
    brs = sum(len(c["branches"]) for c in categories)
    subs = sum(len(b["sub_branches"]) for c in categories for b in c["branches"])
    return secs, brs, subs


def items_outside_tree(categories, items):
    """Unique category paths used by items that do NOT exist in the given tree."""
    ids = {path_id(p) for p in all_paths(categories)}
    bad, seen = [], set()
    for it in items:
        path = tuple(x for x in (it["section"], it.get("branch", ""), it.get("sub_branch", "")) if x)
        pid = path_id(path)
        if pid not in ids and pid not in seen:
            seen.add(pid); bad.append(path)
    return bad


def canonicalize_items(items, categories):
    """Give every item the exact spelling of the new tree (e.g. different capital letters)."""
    look = _lookup(categories)
    out = []
    for it in items:
        key = (cat_norm(it["section"]), cat_norm(it.get("branch", "")), cat_norm(it.get("sub_branch", "")))
        s, b, sb = look.get(key, (it["section"], it.get("branch", ""), it.get("sub_branch", "")))
        out.append(dict(it, section=s, branch=b, sub_branch=sb))
    return out


def prune_image_keys(imgmap, categories):
    """Keep only images whose level still exists. -> (new_map, removed_tokens)."""
    ids = {path_id(p) for p in all_paths(categories)}
    keep = {k: v for k, v in imgmap.items() if k in ids}
    return keep, [v for k, v in imgmap.items() if k not in ids and v]
