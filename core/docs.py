"""PDF and Excel documents, generated in memory. All user text is neutralised/escaped."""
import io
import os
import re
from concurrent.futures import ThreadPoolExecutor
from datetime import datetime
from xml.sax.saxutils import escape
from zoneinfo import ZoneInfo

import requests
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

from .validators import display_phone, neutralize_formula

FONT_DIR = os.path.join(os.path.dirname(os.path.dirname(__file__)), "fonts")
_fonts_ready = False


def shape(t):
    """Arabic reshaping + bidi (customer names may be Arabic)."""
    t = str(t)
    if not re.search(r"[\u0600-\u06FF]", t):
        return t
    try:
        import arabic_reshaper
        try:
            from bidi.algorithm import get_display
        except ImportError:
            from bidi import get_display
        return get_display(arabic_reshaper.reshape(t))
    except Exception:
        return t


def local_time(iso, tz):
    try:
        d = datetime.fromisoformat(iso.replace("Z", "+00:00"))
        return d.astimezone(ZoneInfo(tz or "Africa/Tripoli")).strftime("%Y-%m-%d %H:%M")
    except Exception:
        return iso or ""


# ------------------------------------------------------------------ Excel
_HF = Font(name="Arial", bold=True, color="39FF14", size=11)
_HFILL = PatternFill("solid", fgColor="1C2326")
_B = Border(*(Side(style="thin", color="BFC8CC"),) * 4)


def _xlsx(headers, rows, widths, text_cols=()):
    wb = Workbook()
    ws = wb.active
    ws.append(headers)
    for c in ws[1]:
        c.font, c.fill, c.border = _HF, _HFILL, _B
        c.alignment = Alignment(horizontal="center", vertical="center")
    for r in rows:
        ws.append([neutralize_formula(v) if isinstance(v, str) else v for v in r])
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.font, c.border = Font(name="Arial", size=10), _B
            c.alignment = Alignment(vertical="center", wrap_text=True)
            if c.column - 1 in text_cols:
                c.number_format = "@"
    for i, w in enumerate(widths):
        ws.column_dimensions[chr(65 + i)].width = w
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = "A1:%s%d" % (chr(64 + len(headers)), max(ws.max_row, 1))
    b = io.BytesIO()
    wb.save(b)
    return b.getvalue()


def order_xlsx(order):
    """Customer/order copy: EXACTLY Item No | Item Name | Quantity."""
    rows = [(l["item_no"], l["description"], int(l["qty"])) for l in order["lines"]]
    return _xlsx(["Item No", "Item Name", "Quantity"], rows, [22, 60, 11], text_cols=(0,))


def export_xlsx(orders, tz="Africa/Tripoli"):
    rows = []
    for o in orders:
        for l in o["lines"]:
            rows.append((o["order_id"], local_time(o["created_at"], tz), o["customer"]["name"], o["customer"]["phone"],
                         o.get("project_name", ""), l["item_no"], l["description"], int(l["qty"]), l.get("note", ""),
                         o.get("opened_by_display") or ""))
    return _xlsx(["Order ID", "Date", "Customer Name", "Phone", "Project Name", "Item No", "Item Name", "Quantity",
                  "Line Note", "Sales Operator"], rows, [18, 17, 24, 17, 22, 20, 50, 10, 26, 18], text_cols=(0, 3, 5))


def simple_xlsx(headers, rows):
    return _xlsx(headers, rows, [24] * len(headers))


# ------------------------------------------------------------------ PDF
def _register_fonts():
    global _fonts_ready
    if _fonts_ready:
        return
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    pdfmetrics.registerFont(TTFont("DejaVu", os.path.join(FONT_DIR, "DejaVuSans.ttf")))
    pdfmetrics.registerFont(TTFont("DejaVu-Bold", os.path.join(FONT_DIR, "DejaVuSans-Bold.ttf")))
    _fonts_ready = True


def fetch_thumbs(urls, timeout=5):
    """{key: bytes|None} fetched in parallel; failures become placeholders."""
    def one(u):
        try:
            r = requests.get(u, timeout=timeout)
            return r.content if r.status_code == 200 and len(r.content) < 400_000 else None
        except requests.RequestException:
            return None
    with ThreadPoolExecutor(8) as ex:
        keys = list(urls)
        return dict(zip(keys, ex.map(one, [urls[k] for k in keys])))


def order_pdf(order, settings, logo_bytes=None, thumbs=None):
    """A4 PDF. thumbs: {item_no: jpeg bytes or None}."""
    from PIL import Image as PILImage
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.utils import ImageReader
    from reportlab.platypus import Image, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
    _register_fonts()
    thumbs = thumbs or {}
    store = settings.get("store_name", "")
    tz = settings.get("timezone", "Africa/Tripoli")
    support = settings.get("support_number", "")
    hours = settings.get("working_hours_text", "")
    st = ParagraphStyle("b", fontName="DejaVu", fontSize=8.5, leading=11)
    sb = ParagraphStyle("bb", parent=st, fontName="DejaVu-Bold")
    big = ParagraphStyle("big", fontName="DejaVu-Bold", fontSize=12, leading=15)
    P = lambda t, s=st: Paragraph(escape(shape(t)), s)

    def decorate(c, doc):
        w, h = A4
        c.saveState()
        c.setFillColor(colors.HexColor("#1c2326"))
        c.rect(0, h - 62, w, 62, fill=1, stroke=0)
        x = 36
        if logo_bytes:
            try:
                ir = ImageReader(io.BytesIO(logo_bytes))
                iw, ih = ir.getSize()
                hh = 38
                c.drawImage(ir, x, h - 52, width=hh * iw / ih, height=hh, mask="auto")
                x += hh * iw / ih + 10
            except Exception:
                pass
        c.setFillColor(colors.HexColor("#39ff14"))
        c.setFont("DejaVu-Bold", 14)
        c.drawString(x, h - 32, shape(store))
        c.setFillColor(colors.white)
        c.setFont("DejaVu", 9)
        c.drawRightString(w - 36, h - 28, "RFQ  " + order["order_id"])
        c.drawRightString(w - 36, h - 42, local_time(order["created_at"], tz))
        c.setFillColor(colors.HexColor("#4a585d"))
        c.setFont("DejaVu", 7.5)
        foot = shape(store)
        if support:
            foot += "   |   Support: " + support
        if hours:
            foot += "   |   " + shape(hours)[:80]
        c.drawString(36, 22, foot)
        c.drawRightString(w - 36, 22, "Page %d" % doc.page)
        c.restoreState()

    doc = SimpleDocTemplate(io.BytesIO(), pagesize=A4, leftMargin=36, rightMargin=36, topMargin=78, bottomMargin=40,
                            title="RFQ " + order["order_id"])
    buf = io.BytesIO()
    doc.filename = buf
    cu = order["customer"]
    story = [Paragraph("Request for Quotation", big), Spacer(1, 4),
             P("Customer: " + cu["name"], sb), P("Phone: " + display_phone(cu["phone"])),
             P("Sales Operator: " + order["opened_by_display"]) if order.get("opened_by_display") else Spacer(1, 0)]
    if order.get("project_name"):
        story.append(P("Project: " + order["project_name"]))
    story.append(Spacer(1, 8))
    data = [[P("#", sb), P("Photo", sb), P("Item No", sb), P("Description", sb), P("Qty", sb)]]
    for i, l in enumerate(order["lines"], 1):
        img = thumbs.get(l["item_no"])
        cell = P("No photo")
        if img:
            try:
                w, h = PILImage.open(io.BytesIO(img)).size
                sc = min(44 / w, 44 / h)
                cell = Image(io.BytesIO(img), width=w * sc, height=h * sc)
            except Exception:
                pass
        desc = escape(shape(l["description"]))
        if l.get("note"):
            desc += "<br/><i>Note: %s</i>" % escape(shape(l["note"]))
        data.append([P(str(i)), cell, Paragraph(escape(shape(l["item_no"])), sb), Paragraph(desc, st), P(str(int(l["qty"])), sb)])
    t = Table(data, colWidths=[24, 52, 100, 302, 45], repeatRows=1)
    t.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#e5ece9")),
                           ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#9aa7ab")),
                           ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                           ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f6f9f8")])]))
    story.append(t)
    doc.build(story, onFirstPage=decorate, onLaterPages=decorate)
    return buf.getvalue()
